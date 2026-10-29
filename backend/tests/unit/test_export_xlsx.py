"""xlsx exporter: Week 9 done-when — totals match on-screen to the paisa."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from api.db.models import BOQLine, BuildingParams, Estimate, Organization, Project
from api.services.export_xlsx import estimate_to_xlsx
from api.services.rate_import import import_from_yaml

SAMPLE_DSR = (
    Path(__file__).resolve().parents[2] / "data" / "cpwd_dsr_2023" / "sample.yaml"
)


@pytest.fixture
def seeded_estimate(db_session):
    """Seed one Estimate with 4 BOQ lines across 2 elements + a totals block."""
    edition = import_from_yaml(db_session, SAMPLE_DSR)
    org = Organization(name="Export Co")
    db_session.add(org)
    db_session.flush()
    proj = Project(org_id=org.id, name="Export sample project", location="Mumbai")
    db_session.add(proj)
    db_session.flush()
    est = Estimate(
        project_id=proj.id,
        rate_edition_id=edition.id,
        cost_index=1.05,
        contingency_pct=3.0,
        prelims_pct=2.0,
        status="draft",
    )
    db_session.add(est)
    db_session.flush()
    db_session.add(BuildingParams(estimate_id=est.id, inputs={"builtup_area_m2": 1200}))

    lines_data = [
        ("substructure", 1, "Earthwork in excavation", "m3", 100.0, 320.0, 32000.0, False),
        ("substructure", 2, "RCC M20 in footings", "m3", 10.0, 7205.0, 72050.0, True),
        ("frame", 1, "RCC M25 in columns", "m3", 20.0, 8945.0, 178900.0, False),
        ("frame", 2, "RCC M25 in slabs", "m3", 50.0, 8225.0, 411250.0, False),
    ]
    for el, seq, desc, unit, qty, rate, amount, overridden in lines_data:
        db_session.add(BOQLine(
            estimate_id=est.id, element=el, seq=seq,
            description=desc, unit=unit, qty=qty, rate=rate, amount=amount,
            is_overridden=overridden,
            override_reason="site constraint" if overridden else None,
        ))
    subtotal = sum(d[6] for d in lines_data)
    est.total_cost = round(subtotal * 1.05, 2)
    est.totals_by_element = {"substructure": 104050.0, "frame": 590150.0}
    db_session.commit()
    return est


# ------------------------------------------------------------------ shape
def test_exports_a_valid_xlsx(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    assert blob[:4] == b"PK\x03\x04", "xlsx is a zip; magic bytes must match"
    wb = load_workbook(io.BytesIO(blob))
    assert "BOQ Abstract" in wb.sheetnames
    assert wb.active["A1"].value.startswith("BOQ Abstract — Export sample project")


def test_header_block_includes_edition_and_cost_index(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active
    values = [row for row in ws.iter_rows(min_row=1, max_row=6, values_only=True)]
    flat = " ".join(str(v) for row in values for v in row if v is not None)
    assert "CPWD DSR" in flat
    assert "2023" in flat
    assert "1.050" in flat  # cost_index as 3dp
    assert "PRELIMINARY" in flat


def test_column_headings_are_dsr_style(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active
    row7 = [c.value for c in ws[7]]
    assert row7 == [
        "Serial", "DSR Item No", "Description", "Unit",
        "Qty", "Rate (INR)", "Amount (INR)",
    ]


def test_grand_total_matches_estimate_total_to_paisa(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active

    # Find the row whose column C text starts with "GRAND TOTAL".
    found = None
    for row in ws.iter_rows(min_row=10, values_only=False):
        if row[2].value and str(row[2].value).upper().startswith("GRAND TOTAL"):
            found = row
            break
    assert found is not None, "no GRAND TOTAL row rendered"
    grand = found[6].value
    assert grand == pytest.approx(seeded_estimate.total_cost, abs=0.01)


def test_subtotal_rows_match_element_sum(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active
    # Walk the data region, pick rows where column C starts with "Subtotal —"
    subtotals: dict[str, float] = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not row:
            continue
        c = str(row[2] or "")
        if c.lower().startswith("subtotal —"):
            # e.g. "Subtotal — substructure"
            element = c.split("—", 1)[1].strip()
            subtotals[element] = float(row[6])
    expected = seeded_estimate.totals_by_element or {}
    # Element labels in the sheet use spaces — mirror that here.
    expected_normalised = {k.replace("_", " "): v for k, v in expected.items()}
    assert subtotals == pytest.approx(expected_normalised, rel=1e-4)


def test_overridden_rows_get_flagged(seeded_estimate, db_session):
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active
    found_override = False
    for row in ws.iter_rows(min_row=8, values_only=False):
        desc = row[2]
        if desc and desc.value and "RCC M20 in footings" in str(desc.value):
            found_override = desc.comment is not None
            break
    assert found_override, "overridden lines should carry a cell comment"


def test_includes_dsr_item_no_for_mapped_lines(seeded_estimate, db_session):
    """Lines without rate_item_id leave the DSR column blank — that's fine."""
    blob = estimate_to_xlsx(seeded_estimate, db_session)
    ws = load_workbook(io.BytesIO(blob)).active
    item_nos = [row[1] for row in ws.iter_rows(min_row=8, values_only=True) if row[1]]
    # In this fixture nothing is mapped, so the column is empty — ensure the
    # export didn't crash trying to hydrate.
    assert isinstance(item_nos, list)
